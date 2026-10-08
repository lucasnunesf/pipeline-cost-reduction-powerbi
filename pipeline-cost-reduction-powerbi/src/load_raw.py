"""
load_raw.py

Reads every supplier workbook and the finance targets file and lands them
in SQLite staging tables, exactly as they were typed.

  stg_supplier_ideas   one row per idea row found in a supplier file
  stg_targets          one row per vehicle and month of the targets file
  stg_load_log         one row per file read, with how many rows it gave

Rules of this step:
  - values are stored as TEXT, as they were typed. Nothing is fixed here.
    ("12,40" stays "12,40". Cleaning is the next step.)
  - every row keeps the file and the row number it came from.
  - the template has 200 rows with formulas already in them. A row counts as
    an idea row only if at least one input column was typed; the rest are
    the template's empty rows.
  - every run rebuilds the staging tables from zero (full reload).

Run from the repository root:
    python src/load_raw.py
"""

import sqlite3
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
DB = ROOT / "data" / "cost_reduction.db"

FIRST_ROW = 8   # first data row in the template

# template column -> staging column. Calculated columns K and L are not
# loaded: they are recalculated later. N is loaded only to detect overwrites.
COLUMNS = {
    "A": "idea_id", "B": "proposal_name", "C": "description",
    "D": "responsible", "E": "buyer", "F": "areas", "G": "vehicle",
    "H": "part_number", "I": "current_cost", "J": "new_cost",
    "M": "annual_volume", "N": "annual_saving_in_file", "O": "status",
    "P": "submission_date", "Q": "planned_date", "R": "actual_date",
    "S": "comments",
}
INPUT_COLUMNS = [c for c in COLUMNS if c != "N"]


def as_text(value):
    """Store any cell value as text, without fixing it."""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")   # a real Excel date: keep it ISO
    return str(value)


def col_index(letter):
    return ord(letter) - ord("A")


def create_tables(con):
    cols = ",\n    ".join(f"{name} TEXT" for name in COLUMNS.values())
    con.executescript(f"""
    DROP TABLE IF EXISTS stg_supplier_ideas;
    DROP TABLE IF EXISTS stg_targets;
    DROP TABLE IF EXISTS stg_load_log;

    CREATE TABLE stg_supplier_ideas (
        source_file      TEXT NOT NULL,
        source_row       INTEGER NOT NULL,
        supplier_raw     TEXT,
        supplier_code    TEXT,
        commodity_group  TEXT,
        last_update_raw  TEXT,
        {cols}
    );

    CREATE TABLE stg_targets (
        source_file  TEXT NOT NULL,
        vehicle_raw  TEXT,
        month_raw    TEXT,
        target_raw   TEXT
    );

    CREATE TABLE stg_load_log (
        source_file  TEXT NOT NULL,
        rows_loaded  INTEGER NOT NULL,
        loaded_at    TEXT NOT NULL
    );
    """)


def load_supplier_file(con, path):
    wb = load_workbook(path, read_only=True)
    ws = wb["Proposals"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    wb.close()

    header = {
        "supplier_raw": as_text(rows[0][1]),     # B3
        "supplier_code": as_text(rows[0][4]),    # E3
        "commodity_group": as_text(rows[1][1]),  # B4
        "last_update_raw": as_text(rows[1][4]),  # E4
    }
    source = str(path.relative_to(RAW))

    records = []
    for offset, row in enumerate(rows[FIRST_ROW - 3:]):
        row = list(row) + [None] * (19 - len(row))
        if all(row[col_index(c)] in (None, "") for c in INPUT_COLUMNS):
            continue                              # nothing typed in this row
        record = {"source_file": source, "source_row": FIRST_ROW + offset, **header}
        for letter, name in COLUMNS.items():
            record[name] = as_text(row[col_index(letter)])
        records.append(record)

    if records:
        names = list(records[0])
        con.executemany(
            f"INSERT INTO stg_supplier_ideas ({', '.join(names)}) "
            f"VALUES ({', '.join('?' for _ in names)})",
            [tuple(r[n] for n in names) for r in records],
        )
    return source, len(records)


def load_targets(con, path):
    """Finance sends one column per month. Land it as one row per
    vehicle and month (same values, only the shape changes)."""
    wb = load_workbook(path, read_only=True)
    rows = [r for r in wb["Targets"].iter_rows(values_only=True)]
    wb.close()

    header_i = next(i for i, r in enumerate(rows) if r and r[0] == "Vehicle")
    months = rows[header_i][1:]
    records = []
    for r in rows[header_i + 1:]:
        if not r or r[0] is None:
            continue
        for month, value in zip(months, r[1:]):
            records.append((path.name, as_text(r[0]), as_text(month), as_text(value)))

    con.executemany("INSERT INTO stg_targets VALUES (?, ?, ?, ?)", records)
    return path.name, len(records)


def main():
    files = sorted(RAW.rglob("Supplier*.xlsx"))
    if not files:
        raise SystemExit("No supplier files found. Run src/generate_raw.py first.")

    con = sqlite3.connect(DB)
    create_tables(con)
    now = datetime.now().isoformat(timespec="seconds")

    log = [load_supplier_file(con, f) for f in files]
    log.append(load_targets(con, RAW / "targets.xlsx"))
    con.executemany("INSERT INTO stg_load_log VALUES (?, ?, ?)",
                    [(src, n, now) for src, n in log])
    con.commit()

    ideas = con.execute("SELECT COUNT(*) FROM stg_supplier_ideas").fetchone()[0]
    empty = con.execute(
        "SELECT COUNT(*) FROM stg_load_log WHERE rows_loaded = 0").fetchone()[0]
    targets = con.execute("SELECT COUNT(*) FROM stg_targets").fetchone()[0]
    con.close()
    print(f"{len(files)} supplier files read ({empty} with no ideas), "
          f"{ideas} idea rows, {targets} target rows -> {DB.name}")


if __name__ == "__main__":
    main()
